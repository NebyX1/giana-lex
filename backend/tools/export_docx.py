import re
from io import BytesIO

from docx import Document


def export_markdown_docx(markdown: str) -> BytesIO:
    document = Document()
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph()
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
            continue
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
