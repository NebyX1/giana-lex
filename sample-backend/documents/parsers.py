from __future__ import annotations

import re
import unicodedata
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Callable


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
DEFAULT_MAX_BYTES = 20 * 1024 * 1024


class DocumentParseError(ValueError):
    """A document could not be read safely or its optional parser is absent."""


@dataclass(frozen=True)
class ParsedDocument:
    path: str
    filename: str
    format: str
    text: str
    size_bytes: int

    def to_dict(self) -> dict[str, object]:
        return asdict(self)


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text.replace("\r\n", "\n").replace("\r", "\n"))
    text = text.replace("\x00", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{4,}", "\n\n\n", text).strip()


def resolve_safe_path(path: str | Path, allowed_root: str | Path | None = None) -> Path:
    raw = Path(path).expanduser()
    if allowed_root is not None and not raw.is_absolute():
        raw = Path(allowed_root).expanduser() / raw
    candidate = raw.resolve(strict=True)
    if allowed_root is not None:
        root = Path(allowed_root).expanduser().resolve(strict=True)
        try:
            candidate.relative_to(root)
        except ValueError as exc:
            raise DocumentParseError("La ruta esta fuera del directorio permitido.") from exc
    if not candidate.is_file():
        raise DocumentParseError("La ruta no corresponde a un archivo regular.")
    return candidate


def _read_plain(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("El archivo de texto debe usar codificacion UTF-8.") from exc


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentParseError(
            "No se puede leer PDF: instale la dependencia opcional 'pypdf'."
        ) from exc
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise DocumentParseError("No se admiten PDF cifrados.")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"No se pudo procesar el PDF: {exc}") from exc


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParseError(
            "No se puede leer DOCX: instale la dependencia opcional 'python-docx'."
        ) from exc
    try:
        document = Document(str(path))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(blocks)
    except Exception as exc:
        raise DocumentParseError(f"No se pudo procesar el DOCX: {exc}") from exc


READERS: dict[str, Callable[[Path], str]] = {
    ".txt": _read_plain,
    ".md": _read_plain,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
}


def parse_document(
    path: str | Path,
    *,
    allowed_root: str | Path | None = None,
    max_bytes: int = DEFAULT_MAX_BYTES,
) -> ParsedDocument:
    if max_bytes <= 0:
        raise DocumentParseError("El limite de tamano debe ser positivo.")
    resolved = resolve_safe_path(path, allowed_root)
    extension = resolved.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            f"Formato no admitido '{extension or 'sin extension'}'. Use TXT, MD, PDF o DOCX."
        )
    size = resolved.stat().st_size
    if size > max_bytes:
        raise DocumentParseError(f"El archivo supera el limite de {max_bytes} bytes.")
    text = normalize_text(READERS[extension](resolved))
    return ParsedDocument(str(resolved), resolved.name, extension[1:], text, size)
