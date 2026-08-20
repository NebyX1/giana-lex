from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

from documents.parsers import DocumentParseError
from documents.resolution import LegalResolutionDraft
from .base import Tool


def _safe_docx_output(output_path: str | Path, allowed_root: str | Path) -> Path:
    root = Path(allowed_root).expanduser().resolve(strict=True)
    raw_output = Path(output_path).expanduser()
    if not raw_output.is_absolute():
        raw_output = root / raw_output
    output = raw_output.resolve(strict=False)
    try:
        output.relative_to(root)
    except ValueError as exc:
        raise DocumentParseError("La salida esta fuera del directorio permitido.") from exc
    if output.suffix.lower() != ".docx":
        raise DocumentParseError("La salida debe tener extension .docx.")
    if not output.parent.is_dir():
        raise DocumentParseError("El directorio de salida no existe.")
    return output


def render_text_docx(content: str, output_path: str | Path, allowed_root: str | Path) -> dict[str, object]:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParseError(
            "No se puede crear DOCX: instale la dependencia opcional 'python-docx'."
        ) from exc

    output = _safe_docx_output(output_path, allowed_root)
    document = Document()
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph()
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            document.add_heading(re.sub(r"\*\*(.+?)\*\*", r"\1", heading.group(2)), level=min(len(heading.group(1)), 3))
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            document.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", bullet.group(1)), style="List Bullet")
            continue
        document.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", line))
    document.save(str(output))
    return {"path": str(output), "format": "docx", "validated": True, "legal_approval": False}


class CreateDocxInput(BaseModel):
    resolution: LegalResolutionDraft
    output_path: str = Field(min_length=1, max_length=4096)


def render_resolution_docx(resolution: LegalResolutionDraft, output_path: str | Path, allowed_root: str | Path) -> dict[str, object]:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise DocumentParseError(
            "No se puede crear DOCX: instale la dependencia opcional 'python-docx'."
        ) from exc

    output = _safe_docx_output(output_path, allowed_root)

    document = Document()
    title = document.add_heading(resolution.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if resolution.number:
        paragraph = document.add_paragraph(resolution.number)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if resolution.place or resolution.date:
        document.add_paragraph(", ".join(value for value in (resolution.place, resolution.date) if value))
    if resolution.visto:
        document.add_heading("VISTO", level=1)
        document.add_paragraph(resolution.visto)
    if resolution.resultandos:
        document.add_heading("RESULTANDO", level=1)
        for index, value in enumerate(resolution.resultandos, 1):
            document.add_paragraph(f"{index}) {value}")
    if resolution.considerandos:
        document.add_heading("CONSIDERANDO", level=1)
        for index, value in enumerate(resolution.considerandos, 1):
            document.add_paragraph(f"{index}) {value}")
    if resolution.atento:
        document.add_heading("ATENTO", level=1)
        document.add_paragraph(resolution.atento)
    document.add_heading("RESUELVE", level=1)
    for article in resolution.resuelve:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"ARTICULO {article.number}: ").bold = True
        paragraph.add_run(article.text)
    if resolution.notes:
        document.add_heading("NOTAS DE REVISION", level=1)
        for value in resolution.notes:
            document.add_paragraph(value)
    document.save(str(output))
    return {"path": str(output), "format": "docx", "validated": True, "legal_approval": False}


class CreateDocx(Tool[CreateDocxInput]):
    name = "create_docx"
    description = "Renderiza JSON de resolucion validado a DOCX; no implica aprobacion juridica."
    input_schema = CreateDocxInput

    def execute(self, arguments: CreateDocxInput, context: dict[str, Any]) -> dict[str, object]:
        return render_resolution_docx(
            arguments.resolution,
            arguments.output_path,
            context.get("document_root", Path.cwd()),
        )


tool = CreateDocx()
